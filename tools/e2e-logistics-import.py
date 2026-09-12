# -*- coding: utf-8 -*-
"""物流公司「基本信息导入」接口端到端验证（独立后端 5662）"""
import base64, io, json, re, sys, urllib.request, uuid, time
from urllib.parse import quote

BASE = "http://localhost:5662/api"
results = []

def check(name, ok, detail=""):
    results.append((name, ok, detail))
    print(("PASS " if ok else "FAIL ") + name + (" — " + str(detail) if detail else ""))

def call(path, data=None, method=None, token=None, raw=False, files=None):
    url = BASE + path
    headers = {}
    if files:
        boundary = "----e2e" + uuid.uuid4().hex
        body = b""
        for k, v in (data or {}).items():
            body += ("--%s\r\nContent-Disposition: form-data; name=\"%s\"\r\n\r\n%s\r\n" % (boundary, k, v)).encode()
        fname, content = files
        body += ("--%s\r\nContent-Disposition: form-data; name=\"file\"; filename=\"%s\"\r\n"
                 "Content-Type: application/vnd.openxmlformats-officedocument.spreadsheetml.sheet\r\n\r\n" % (boundary, fname)).encode()
        body += content + ("\r\n--%s--\r\n" % boundary).encode()
        headers["Content-Type"] = "multipart/form-data; boundary=" + boundary
    elif data is not None:
        body = json.dumps(data).encode()
        headers["Content-Type"] = "application/json"
    else:
        body = None
    if token:
        headers["Authorization"] = "Bearer " + token
    req = urllib.request.Request(url, data=body, headers=headers, method=method or ("POST" if body else "GET"))
    try:
        with urllib.request.urlopen(req, timeout=60) as resp:
            b = resp.read()
            return (b, resp.status) if raw else json.loads(b)
    except urllib.error.HTTPError as e:
        b = e.read()
        return (b, e.code) if raw else {"__http": e.code, "__body": b.decode("utf-8", "ignore")[:300]}

def login():
    cap = call("/auth/captcha")
    svg = base64.b64decode(cap["data"]["img"].split(",", 1)[1]).decode("utf-8", "ignore")
    code = "".join(re.findall(r"<text[^>]*>([^<]+)</text>", svg))
    res = call("/auth/login", {"username": "admin", "password": "admin123", "tenantName": "系统租户",
                               "captcha": code, "captchaKey": cap["data"]["uuid"]})
    tok = (res.get("data") or {}).get("token")
    if not tok:
        raise SystemExit("login failed: " + json.dumps(res, ensure_ascii=False)[:300])
    return tok

TOKEN = login()
print("login ok")

# ── 1. 模板下载 ──
blob, status = call("/erp/md/customer/import-template?partnerType=logistics", token=TOKEN, raw=True)
check("模板下载返回 200", status == 200, "status=%s" % status)
check("模板为真实 xlsx（PK 魔数）", blob[:2] == b"PK", blob[:2])
open("I:/AI-Ready/tool-results/ql361/logistics-shots/ours-logistics-import-template.xlsx", "wb").write(blob)

import openpyxl
wb = openpyxl.load_workbook(io.BytesIO(blob))
ws = wb.worksheets[0]
headers = [c.value for c in ws[1]]
check("模板表头与对标一致", headers == ["导入结果", "物流公司编号(必填)", "物流公司名称(必填)", "联系人", "联系电话", "公司地址", "备注"],
      json.dumps(headers, ensure_ascii=False))
check("模板首行含批注说明", any(c.comment for c in ws[1]),
      ",".join((c.comment.text[:12] if c.comment else "") for c in ws[1]))

# ── 2. 导入 Excel（真实落库） ──
stamp = time.strftime("%m%d%H%M%S")
wb2 = openpyxl.Workbook()
ws2 = wb2.active
ws2.append(headers)
ws2.append(["", "E2EWL" + stamp, "E2E物流导入甲公司" + stamp, "赵六", "13800000088", "酒泉市导入路8号", "导入备注A"])
ws2.append(["", "", "", "", "", "", ""])                       # 空行 → 跳过
ws2.append(["", "", "E2E缺编号名称公司" + stamp, "钱七", "13800000099", "酒泉市导入路9号", "自动编号"])
buf = io.BytesIO(); wb2.save(buf)
res = call("/erp/md/customer/import-excel?partnerType=logistics", data={}, token=TOKEN, files=("e2e.xlsx", buf.getvalue()))
data = res.get("data") or {}
check("导入接口返回成功统计", data.get("total") == 2 and data.get("success") == 2, json.dumps(data, ensure_ascii=False)[:200])

# ── 3. 回读校验 ──
page = call("/erp/md/customer/page?partnerType=LOGISTICS&pageNum=1&pageSize=10&keyword=" + quote("E2E物流导入甲公司" + stamp), token=TOKEN)
rows = (page.get("data") or {}).get("records") or []
check("导入记录已落库（按名称检索）", len(rows) == 1, "total=%s" % (page.get("data") or {}).get("total"))
row = rows[0] if rows else {}
check("导入带出联系人（主网点 linkman）", row.get("contactPerson") == "赵六", str(row.get("contactPerson")))
check("导入带出联系电话", row.get("contactPhone") == "13800000088", str(row.get("contactPhone")))
check("导入带出公司地址", row.get("address") == "酒泉市导入路8号", str(row.get("address")))
check("导入备注落库", row.get("remark") == "导入备注A", str(row.get("remark")))
check("导入记录类型为物流公司", row.get("partnerType") == "logistics", str(row.get("partnerType")))
check("显式编号原样落库", row.get("partnerCode") == "E2EWL" + stamp, str(row.get("partnerCode")))

page2 = call("/erp/md/customer/page?partnerType=LOGISTICS&pageNum=1&pageSize=10&keyword=" + quote("E2E缺编号名称公司" + stamp), token=TOKEN)
rows2 = (page2.get("data") or {}).get("records") or []
code2 = rows2[0].get("partnerCode") if rows2 else None
check("缺编号自动按 WuLiu 口径补号", bool(code2) and re.match(r"^WuLiu\d{3,}$", code2 or ""), str(code2))

# ── 4. 重号拒绝 ──
buf3 = io.BytesIO(); wb3 = openpyxl.Workbook(); ws3 = wb3.active
ws3.append(headers); ws3.append(["", "E2EWL" + stamp, "重号公司", "", "", "", ""]); wb3.save(buf3)
res3 = call("/erp/md/customer/import-excel?partnerType=logistics", data={}, token=TOKEN, files=("dup.xlsx", buf3.getvalue()))
d3 = res3.get("data") or {}
check("编号重复被拒绝并给出原因", d3.get("failure") == 1 and any("已存在" in e for e in (d3.get("errors") or [])),
      json.dumps(d3, ensure_ascii=False)[:200])

# ── 5. 清理 ──
for r in rows + rows2:
    call("/erp/md/customer/" + str(r["id"]), method="DELETE", token=TOKEN)

ok = sum(1 for _, o, _ in results if o)
print("\n==== %d/%d PASS ====" % (ok, len(results)))
sys.exit(0 if ok == len(results) else 1)
